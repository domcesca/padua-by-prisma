"""CMS hospital wage indexes: each California hospital's wage index for IPPS (inpatient) and OPPS
(outpatient) payments, for Propose's Advanced-mode wage index adjustment.

Medicare splits a DRG or APC payment into a labor-related share, multiplied by the hospital's
area wage index, and a non-labor share, which isn't. Propose's standard estimate is the flat
national rate; Advanced mode can apply the hospital's own wage index the way CMS does.

Sources:
  * IPPS: the newest final rule's "Table 2: Case-Mix Index and Wage Index Table by CCN" (in the
    Tables 2, 3, 4A, 4B zip). The payment wage index is the "FY <year> Wage Index With Cap"
    column (footnote 3: "the hospital's FY <year> wage index with all applicable adjustments":
    reclassification, rural and imputed floors, frontier, out-migration, the 5% cap), or the
    low-wage-index transition column where CMS filled it in. Checked against the same rule's
    Impact File, whose wage index must agree for every hospital in both. The labor-related
    shares come from Tables 1A (wage index above 1) and 1B (at or below 1).
  * OPPS: the calendar-year OPPS final rule's Hospital Impact File ("facility-specific impacts"),
    column "Post Reclassification Wage Index": per CMS's file layout, the final FY <year> IPPS
    wage index with the same adjustments, which is what OPPS pays on for the calendar year.
    Paired with the year of cms-opps's Addendum A. Checked against Table 2's "FY <year> Wage
    Index" column. OPPS's labor-related share is 60% (unchanged since the OPPS began; OPPS
    final rule, section II.C); the file doesn't restate it.

Children's and cancer hospitals aren't in the OPPS Impact File, and critical access, children's,
cancer, psychiatric, rehabilitation, and long-term care hospitals aren't in Table 2 (they aren't
paid under IPPS), so those hospitals have no wage index here and Propose keeps the national rate.
Run after cms-opps.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile

import openpyxl

from ..core import PROCESSED_DIR, Dataset, Resource, utc_now_iso, write_json
from ..crosswalk import FacilityCrosswalk, match_ccns
from .cms_ipps import CMS, IPPS_PAGE, _get, _xlsx, parse_table1
from .cms_penalties import GENERAL_ID, _num, _pdc_csv

OPPS_RULES = f"{CMS}/medicare/payment/prospective-payment-systems/hospital-outpatient/regulations-notices"
OPPS_LABOR_SHARE = 0.60
STATE_CODE = "05"  # California's CCN prefix


def _zip_member(path, pattern: str) -> bytes:
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if re.search(pattern, n, re.I) and not n.startswith("__MACOSX")]
        if len(names) != 1:
            raise RuntimeError(f"{path.name}: expected one file matching {pattern!r}, found {names}")
        return z.read(names[0])


def _ccn(v) -> str:
    return str(v).strip().zfill(6) if v is not None else ""


class CmsWageIndex(Dataset):
    id = "cms-wage-index"
    title = "CMS hospital wage indexes for IPPS and OPPS payments"
    source_page = IPPS_PAGE

    def resources(self) -> list[Resource]:
        years = sorted({int(y) for y in re.findall(r"/fy-(\d{4})-ipps-final-rule-home-page", _get(IPPS_PAGE))})
        fy = self.years[-1] if self.years else years[-1]
        self.fiscal_year = fy
        self.rule_page = f"{IPPS_PAGE}/fy-{fy}-ipps-final-rule-home-page"
        links = sorted(set(re.findall(r'href="(/files/zip/[^"]+\.zip)"', _get(self.rule_page))))
        wanted = {
            "table2": rf"fy-?{fy}-ipps-fr-tables-2-3-4a-4b\.zip$",
            "table1": rf"fy-?{fy}-ipps-fr-table-1a-1e\.zip$",
            "impact": rf"fy-?{fy}-ipps-fr-impact-file\.zip$",
        }
        out = []
        for key, pattern in wanted.items():
            hit = [link for link in links if re.search(pattern, link)]
            if len(hit) != 1:
                raise RuntimeError(f"FY {fy}: expected one {key} zip on {self.rule_page}, found {hit}")
            out.append(Resource(f"FY {fy} IPPS final rule {key}", CMS + hit[0], "ZIP", tags={"key": key}))

        # The OPPS year: whatever Addendum A cms-opps built from.
        opps_manifest = json.loads((PROCESSED_DIR / "cms-opps" / "manifest.json").read_text())
        cy = self.calendar_year = int(opps_manifest["calendarYear"])
        rules = sorted(set(re.findall(r'href="([^"]*/cms-\d{4}-fc?)"', _get(OPPS_RULES))), reverse=True)
        impact = None
        for rule in rules:
            page = _get(rule if rule.startswith("http") else CMS + rule)
            hit = re.findall(rf'href="(/files/zip/{cy}-nfrm-opps-facility-specific-impacts\.zip)"', page)
            if hit:
                impact, self.opps_rule_page = hit[0], rule if rule.startswith("http") else CMS + rule
                break
        if not impact:
            raise RuntimeError(f"No CY {cy} OPPS final rule facility-specific impact file under {OPPS_RULES}")
        out.append(Resource(f"CY {cy} OPPS final rule Hospital Impact File", CMS + impact, "ZIP", tags={"key": "opps"}))
        general_title, general_url = _pdc_csv(GENERAL_ID)
        out.append(Resource(general_title, general_url, "CSV", tags={"key": "general"}))
        return out

    def load(self, files):
        self.crosswalk = FacilityCrosswalk.load(self.refresh)
        paths = {res.tags["key"]: path for res, path in files}
        self.sources = [{"name": res.name, "url": res.url} for res, _ in files]
        with open(paths["general"], encoding="latin1", newline="") as f:
            general = [r for r in csv.DictReader(f) if r.get("State") == "CA"]
        return {
            "table2": self._table2(paths["table2"]),
            "impact": self._ipps_impact(paths["impact"]),
            "rates": parse_table1(_xlsx(paths["table1"])),
            "opps": self._opps_impact(paths["opps"]),
            "general": general,
        }

    def _table2(self, path) -> dict[str, dict]:
        fy = self.fiscal_year
        text = _zip_member(path, r"table 2\.txt$").decode("latin1")
        rows = list(csv.reader(io.StringIO(text), delimiter="\t"))
        header_at = next((i for i, r in enumerate(rows[:10]) if r and re.fullmatch(r"\d?CCN", r[0].strip())), None)
        if header_at is None:
            raise RuntimeError("Table 2: no CCN header row; the layout changed.")
        header = [re.sub(r"\s+", " ", c).strip() for c in rows[header_at]]

        def col(pattern: str) -> int:
            hits = [i for i, h in enumerate(header) if re.search(pattern, h, re.I)]
            if len(hits) != 1:
                raise RuntimeError(f"Table 2: expected one column matching {pattern!r}, found {[header[i] for i in hits]}")
            return hits[0]

        c = {
            "final": col(rf"FY {fy} Wage Index With Cap$"),
            "transition": col(rf"FY {fy} Transition for the Discontinuation"),
            "prior": col(rf"^FY {fy - 1} Wage Index$"),
            "geoCbsa": col(r"^Geographic CBSA$"),
            "payCbsa": col(r"Wage Index Payment CBSA$"),
        }
        out = {}
        for r in rows[header_at + 1 :]:
            if not r or not re.fullmatch(r"\d{2}[0-9A-Z]\d{3}", r[0].strip()):
                continue
            final = _num(r[c["transition"]]) or _num(r[c["final"]])
            if final is None:
                continue
            out[r[0].strip()] = {
                "wageIndex": final,
                "priorYear": _num(r[c["prior"]]),
                "cbsa": (r[c["payCbsa"]].strip() or r[c["geoCbsa"]].strip()),
                "geographicCbsa": r[c["geoCbsa"]].strip(),
                "reclassified": bool(r[c["payCbsa"]].strip()),
            }
        if len(out) < 3000:
            raise RuntimeError(f"Table 2: only {len(out)} hospitals parsed; expected ~3,200.")
        return out

    def _ipps_impact(self, path) -> dict[str, float]:
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
            if len(names) != 1:
                raise RuntimeError(f"{path.name}: expected one .xlsx, found {names}")
            wb = openpyxl.load_workbook(io.BytesIO(z.read(names[0])), read_only=True, data_only=True)
        sheet = next(ws for ws in wb.worksheets if "variable" not in ws.title.lower())
        rows = [tuple(r) for r in sheet.iter_rows(values_only=True)]
        header_at = next(i for i, r in enumerate(rows[:10]) if str(r[0]).strip() == "Provider Number")
        header = [str(h).strip() for h in rows[header_at]]
        wi = header.index(f"FY {self.fiscal_year} Wage Index")
        return {_ccn(r[0]): v for r in rows[header_at + 1 :] if r and r[0] and (v := _num(r[wi])) is not None}

    def _opps_impact(self, path) -> dict[str, dict]:
        data = _zip_member(path, r"impact file[^/]*\.xlsx$")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        header_at = next((i for i, r in enumerate(rows[:10]) if str(r[0]).strip() == "Provider Number"), None)
        if header_at is None:
            raise RuntimeError("OPPS Impact File: no 'Provider Number' header; the layout changed.")
        header = [str(h).strip() for h in rows[header_at]]
        if "Post Reclassification Wage Index" not in header:
            raise RuntimeError(f"OPPS Impact File: no 'Post Reclassification Wage Index' column in {header}")
        wi, name = header.index("Post Reclassification Wage Index"), header.index("Provider Name")
        return {
            _ccn(r[0]): {"wageIndex": v, "name": str(r[name] or "").strip()}
            for r in rows[header_at + 1 :]
            if r and r[0] and (v := _num(r[wi])) is not None
        }

    def build(self, data: dict) -> None:
        fy, cy = self.fiscal_year, self.calendar_year
        table2, impact, opps, rates = data["table2"], data["impact"], data["opps"], data["rates"]

        # Check 1: Table 2's payment wage index is the Impact File's, for every hospital in both.
        both = [c for c in table2 if c in impact]
        off = [(c, table2[c]["wageIndex"], impact[c]) for c in both if abs(table2[c]["wageIndex"] - impact[c]) > 5e-5]
        if off or len(both) < 3000:
            raise RuntimeError(f"Table 2 vs Impact File: {len(off)} of {len(both)} wage indexes disagree, e.g. {off[:5]}")
        print(f"  IPPS FY {fy}: Table 2 matches the Impact File for all {len(both)} hospitals in both")

        # Check 2: OPPS's wage index is the FY <cy> IPPS index (Table 2's prior-year column when cy = fy - 1).
        opps_check = None
        if cy == fy - 1:
            both = [c for c in opps if c in table2 and table2[c]["priorYear"] is not None]
            off = [c for c in both if abs(opps[c]["wageIndex"] - table2[c]["priorYear"]) > 5e-5]
            if len(off) > 0.02 * len(both):
                raise RuntimeError(f"OPPS Impact File vs Table 2 FY {cy}: {len(off)} of {len(both)} disagree, e.g. {off[:5]}")
            opps_check = {"compared": len(both), "differ": len(off), "differCcns": off}
            print(f"  OPPS CY {cy}: {len(both) - len(off)} of {len(both)} match Table 2's FY {cy} wage index (the rest were revised after the OPPS rule)")

        names = {_ccn(r["Facility ID"]): r["Facility Name"].strip() for r in data["general"]}
        zips = {_ccn(r["Facility ID"]): r["ZIP Code"].strip()[:5].zfill(5) for r in data["general"]}
        ccns = sorted(c for c in set(table2) | set(opps) if c.startswith(STATE_CODE))
        for c in ccns:
            names.setdefault(c, opps.get(c, {}).get("name", ""))
        ccn_to_hcai, shared, by_name = match_ccns(self.crosswalk, ccns, names, zips)

        hospitals: dict[str, dict] = {}
        for c in ccns:
            hcai = ccn_to_hcai.get(c)
            if not hcai:
                continue
            ipps = table2.get(c)
            entry = {
                "ccn": c,
                "cmsName": names.get(c),
                "ipps": {"wageIndex": round(ipps["wageIndex"], 4), "cbsa": ipps["cbsa"], "reclassified": ipps["reclassified"]} if ipps else None,
                "opps": {"wageIndex": round(opps[c]["wageIndex"], 4)} if c in opps else None,
            }
            # Two CCNs on one hospital: keep the one with an IPPS wage index (the current certification).
            if hcai in hospitals and hospitals[hcai]["ipps"] and not entry["ipps"]:
                continue
            hospitals[hcai] = entry

        unmatched = [{"ccn": c, "name": names.get(c)} for c in ccns if c not in ccn_to_hcai]
        write_json(self.out_dir / "hospitals.json", hospitals)
        write_json(
            self.out_dir / "manifest.json",
            {
                "id": self.id,
                "title": self.title,
                "ipps": {
                    "fiscalYear": fy,
                    "sourcePage": self.rule_page,
                    "table": f"FY {fy} IPPS Final Rule, Table 2 (wage index by CCN)",
                    # Labor-related and non-labor parts of the national operating standardized amount.
                    "standardizedAmount": {
                        "wageIndexAboveOne": {k: rates["operating"][k] for k in ("laborRelated", "nonlaborRelated")},
                        "wageIndexAtMostOne": rates["operatingLowWage"],
                    },
                },
                "opps": {
                    "calendarYear": cy,
                    "sourcePage": self.opps_rule_page,
                    "table": f"CY {cy} OPPS Final Rule, Hospital Impact File (post-reclassification wage index)",
                    "laborShare": OPPS_LABOR_SHARE,
                    "checkAgainstTable2": opps_check,
                },
                "sources": self.sources,
                "crosswalk": self.crosswalk.sources,
                "sharedReporting": shared,
                "coverage": {
                    "cmsHospitals": len(ccns),
                    "matchedHospitals": len(hospitals),
                    "withIpps": sum(1 for h in hospitals.values() if h["ipps"]),
                    "withOpps": sum(1 for h in hospitals.values() if h["opps"]),
                    "matchedByNameAndZip": by_name,
                    "unmatched": unmatched,
                },
                "generatedAt": utc_now_iso(),
                "notes": [
                    f"IPPS: Table 2's FY {fy} wage index with all adjustments (or the low-wage-index transition value), checked against the FY {fy} Impact File for every hospital.",
                    f"Wage-adjusted operating payment = relative weight x (labor-related amount x wage index + non-labor amount); Table 1A's split applies above a wage index of 1, Table 1B's at or below.",
                    f"OPPS: the CY {cy} OPPS Hospital Impact File's post-reclassification wage index (the final FY {cy} IPPS wage index). Wage-adjusted rate = national rate x ({OPPS_LABOR_SHARE} x wage index + {1 - OPPS_LABOR_SHARE:.2f}).",
                ],
            },
            compact=False,
        )
        print(f"  {len(hospitals)} California hospitals matched ({len(unmatched)} CCNs unmatched)")
